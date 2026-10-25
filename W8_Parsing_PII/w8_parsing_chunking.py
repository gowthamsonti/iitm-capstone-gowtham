import os
import pymupdf
import numpy
import pdfplumber
import re
from openai import OpenAI
from bs4 import BeautifulSoup
from docx import Document
from pathlib import Path
import numpy as np

SAMPLE_DIR = Path.cwd() / "sample_docs"

SAMPLE_PDF  = SAMPLE_DIR / "company_policies.pdf"
SAMPLE_HTML = SAMPLE_DIR / "product_page.html"
SAMPLE_DOCX = SAMPLE_DIR / "onboarding.docx"

for path in [SAMPLE_PDF,SAMPLE_HTML,SAMPLE_DOCX]:
    assert path.exists(), (
        f"Missing {path.name}. Run: python /generate_sample_docs.py"
    )

print("Sample documents ready:")
for path in [SAMPLE_PDF, SAMPLE_HTML, SAMPLE_DOCX]:
    print(f"  {path.name:30s}  {path.stat().st_size:>6d} bytes")

doc = pymupdf.open(SAMPLE_PDF)
print(f"PDF: {SAMPLE_PDF.name}")
print(f"Pages: {len(doc)}\n")

for page_num, page in enumerate(doc, 1):
    text = page.get_text()
    print(f"══ Page {page_num} ({len(text)} chars) ══")
    # Show first 250 chars of each page
    preview = text[:250].replace('\n', ' | ')
    print(f"{preview}...\n")

doc.close()

#3 pdf plumber on same PDF

with pdfplumber.open(SAMPLE_PDF) as pdf:
    first_page = pdf.pages[0]
    # Text extraction — similar to PyMuPDF for prose
    text = first_page.extract_text()
    print("Text extraction (same as PyMuPDF for prose):")
    print(text[:300], "...\n")
    # Table extraction — pdfplumber's superpower
    tables = first_page.extract_tables()
    print(f"Tables detected on page 1: {len(tables)}")
    for i, table in enumerate(tables, 1):
        print(f"\n══ Table {i} ══")
        for row in table:
            print(f"  {' | '.join(str(c) for c in row)}")

#4 Beautiful soup on html page

html = SAMPLE_HTML.read_text()

# Naive: get all text, no stripping
soup_naive = BeautifulSoup(html, "html.parser")
naive_text = soup_naive.get_text(separator="\n", strip=True)
print("══ NAIVE extraction (nav + footer + scripts included) ══")
print(naive_text[:400])
print(f"\n... {len(naive_text)} chars total\n")

soup_clean = BeautifulSoup(html, "html.parser")
for tag in soup_clean(["nav", "footer", "script", "style"]):
    tag.decompose()

# Get <main> if present, else body
main = soup_clean.find("main") or soup_clean.find("body") or soup_clean
clean_text = main.get_text(separator="\n", strip=True)

print("══ CLEAN extraction (nav + footer + scripts stripped) ══")
print(clean_text[:400])
print(f"\n... {len(clean_text)} chars total")
print(f"\nRemoved: {len(naive_text) - len(clean_text)} chars of noise")

#5 python docx on the DOCX
doc = Document(str(SAMPLE_DOCX))

print(f"DOCX: {SAMPLE_DOCX.name}")
print(f"Paragraphs: {len(doc.paragraphs)}\n")

for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        style = para.style.name
        preview = para.text[:70]
        print(f"  [{i:2d}] {style:15s}  {preview}...")


#6 Scanned PDF

scanned_path = SAMPLE_DIR / "scanned_fake.pdf"
scan_doc = pymupdf.open()
page = scan_doc.new_page()
# Draw a filled rectangle to simulate a scanned image with no text layer
page.draw_rect(pymupdf.Rect(100, 100, 500, 700), fill=(0.9, 0.9, 0.9))
scan_doc.save(str(scanned_path))
scan_doc.close()

# Now try to extract text
doc = pymupdf.open(scanned_path)
extracted = doc[0].get_text()
doc.close()

print(f"Fake 'scanned' PDF: {scanned_path.name}")
print(f"Extracted text length: {len(extracted)} chars")
print(f"Extracted content: {extracted!r}")
print()
print("That's the trap. The parser succeeded — no crash. But zero content.")
print("Your ingestion silently produced 0 chunks from this document.")

# Cleanup — don't clutter sample_docs/
scanned_path.unlink()


#7 Topic 1 Recap
#8 Chunking strategy W6 Baseline
# Load the PDF text (page 1 only, for demo compactness)
doc = pymupdf.open(SAMPLE_PDF)
page1_text = doc[0].get_text()
doc.close()

def chunk_fixed(text, size=400, overlap=40):
    """W6's sliding-window chunker."""
    if len(text) <= size:
        return [text]
    chunks, i = [], 0
    while i < len(text):
        end = min(i + size, len(text))
        chunks.append(text[i:end])
        if end == len(text):
            break
        i = end - overlap
    return chunks

fixed_chunks = chunk_fixed(page1_text, size=400, overlap=40)
print(f"Fixed-size chunking (size=400, overlap=40) → {len(fixed_chunks)} chunks\n")

for i, chunk in enumerate(fixed_chunks):
    ends_mid = not chunk.rstrip().endswith(('.', '!', '?', ':'))
    marker = " ← ends mid-sentence" if ends_mid else ""
    preview = chunk[-60:].replace('\n', ' ')
    print(f"  [{i}] {len(chunk):>4d} chars, ends: '...{preview}'{marker}")

#9 Chunking Strategy - Recursive

def chunk_recursive(text, max_size=400):
    """Split by paragraphs. If a paragraph exceeds max_size, split by sentences."""
    paragraphs = [p.strip() for p in text.split('\n\n') if p.strip()]
    chunks = []
    for para in paragraphs:
        if len(para) <= max_size:
            chunks.append(para)
        else:
            # Sentence-level fallback
            sentences = re.split(r'(?<=[.!?])\s+', para)
            current = ""
            for sent in sentences:
                if len(current) + len(sent) + 1 <= max_size:
                    current = (current + " " + sent).strip()
                else:
                    if current:
                        chunks.append(current)
                    current = sent
            if current:
                chunks.append(current)
    return chunks
recursive_chunks = chunk_recursive(page1_text, max_size=400)
print(f"Recursive chunking (max_size=400) → {len(recursive_chunks)} chunks\n")
for i, chunk in enumerate(recursive_chunks):
    ends_mid = not chunk.rstrip().endswith(('.', '!', '?', ':'))
    marker = " ← ends mid-sentence" if ends_mid else " ← clean end"
    preview = chunk[-60:].replace('\n', ' ')
    print(f"  [{i}] {len(chunk):>4d} chars, ends: '...{preview}'{marker}")

    #10 Chunking strategy - Semantic

assert os.environ.get("OPENAI_API_KEY"), "Set OPENAI_API_KEY for the semantic chunking demo"

client = OpenAI()
# 6 sentences, 2 topics — leave policy + remote work
sentences = [
    "Employees receive 20 days of paid annual leave per year.",
    "Leave requests should be submitted at least two weeks in advance.",
    "Unused leave carries over up to 5 days to the following year.",
    "All employees are eligible for hybrid work arrangements.",
    "The standard expectation is three days in-office per week.",
    "Full-remote status requires VP approval and a business case.",
]

# Embed all 6 in one API call (cheap: ~$0.0001)
resp = client.embeddings.create(
    model="text-embedding-3-small",
    input=sentences,
)
embeddings = [np.array(item.embedding) for item in resp.data]
def cosine(a, b):
    return float(np.dot(a, b) / (np.linalg.norm(a) * np.linalg.norm(b)))

# Show consecutive-sentence similarities
print("Consecutive sentence similarities:\n")
for i in range(len(sentences) - 1):
    sim = cosine(embeddings[i], embeddings[i+1])
    print(f"  s{i} → s{i+1}: cosine = {sim:.3f}")
    print(f"    s{i}: {sentences[i]}")
    print(f"    s{i+1}: {sentences[i+1]}\n")

# Now group consecutive sentences into chunks — split whenever similarity drops below threshold
THRESHOLD = 0.55  # tuned for this demo; real systems calibrate per corpus

chunks = [[sentences[0]]]
for i in range(1, len(sentences)):
    sim = cosine(embeddings[i-1], embeddings[i])
    if sim >= THRESHOLD:
        chunks[-1].append(sentences[i])
    else:
        chunks.append([sentences[i]])  # new chunk starts here

print(f"Semantic chunking (threshold={THRESHOLD}) → {len(chunks)} chunks\n")
for i, chunk_sentences in enumerate(chunks):
    print(f"══ Chunk {i} ({len(chunk_sentences)} sentences) ══")
    for s in chunk_sentences:
        print(f"  • {s}")
    print()

#Chunking strategy - Structure aware

def chunk_docx_structure_aware(path):
    """Split a DOCX into chunks bounded by Heading 1 / Heading 2 styles.
    Attach section_path metadata based on the heading hierarchy."""
    doc = Document(str(path))
    chunks = []
    current_h1 = None
    current_h2 = None
    current_text = []
    def flush():
        if current_text:
            section_path = " > ".join(p for p in [current_h1, current_h2] if p)
            chunks.append({
                "section_path": section_path,
                "text": "\n".join(current_text),
            })
    for para in doc.paragraphs:
        if not para.text.strip():
            continue
        style = para.style.name
        if style == "Title" or style == "Heading 1":
            flush()  # save what we have
            current_text = []
            current_h1 = para.text
            current_h2 = None
        elif style == "Heading 2":
            flush()  # save subsection
            current_text = []
            current_h2 = para.text
        else:
            current_text.append(para.text)
    
    flush()  # final flush
    return chunks

structured = chunk_docx_structure_aware(SAMPLE_DOCX)
print(f"Structure-aware chunking of {SAMPLE_DOCX.name} → {len(structured)} chunks\n")

for i, chunk in enumerate(structured):
    print(f"══ Chunk {i} ══")
    print(f"  section_path: {chunk['section_path']!r}")
    preview = chunk['text'][:120].replace('\n', ' ')
    print(f"  text:         {preview}...")
    print(f"  length:       {len(chunk['text'])} chars\n")
    